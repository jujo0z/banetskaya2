#!/usr/bin/env python3
"""
Тестирование обновлённого документа «Заявление о регистрации по месту ПРЕБЫВАНИЯ».
Backend: https://form-portal-15.preview.emergentagent.com/api
"""
import requests
import io
import sys
from datetime import datetime
from openpyxl import load_workbook

BASE_URL = "https://form-portal-15.preview.emergentagent.com/api"

def test_zayavlenie_fields():
    """TEST 1: GET /api/zayavlenie/fields → проверить наличие новых полей"""
    print("\n" + "="*80)
    print("TEST 1: GET /api/zayavlenie/fields")
    print("="*80)
    
    r = requests.get(f"{BASE_URL}/zayavlenie/fields")
    assert r.status_code == 200, f"Expected 200, got {r.status_code}"
    
    data = r.json()
    assert "keys" in data, "Response must contain 'keys'"
    assert "groups" in data, "Response must contain 'groups'"
    
    keys = data["keys"]
    required_keys = [
        "fio", "birth_year", "doc_name", "passport_series", "passport_number",
        "passport_issued_by", "passport_issue_date", "reg_who", "reg_count",
        "address_locality", "res_street", "res_house", "res_korpus", "res_apartment",
        "stay_term", "from_place", "basis", "sign_date", "area"
    ]
    
    missing = [k for k in required_keys if k not in keys]
    assert not missing, f"Missing required keys: {missing}"
    
    print(f"✅ PASS: Found {len(keys)} keys, all required keys present")
    print(f"   Required keys: {required_keys}")
    return True


def test_zayavlenie_preview_pages():
    """TEST 2: POST /api/zayavlenie/preview → проверить количество страниц"""
    print("\n" + "="*80)
    print("TEST 2: POST /api/zayavlenie/preview (page counts)")
    print("="*80)
    
    # Test 1 record → 2 pages
    payload1 = {
        "records": [{
            "fio": "Иванов Иван Иванович",
            "birth_year": "2006",
            "stay_term": "срок до 30.06.2028",
            "from_place": "г. Лида",
            "basis": "Договор найма № 5 от 21.07.2025",
            "area": "5467,9",
            "sign_date": "21.07.2025"
        }]
    }
    
    r1 = requests.post(f"{BASE_URL}/zayavlenie/preview", json=payload1)
    assert r1.status_code == 200, f"Expected 200, got {r1.status_code}"
    assert r1.headers.get("Content-Type") == "application/pdf", "Expected PDF"
    assert r1.content[:4] == b"%PDF", "Response must be valid PDF"
    
    # Check page count using pymupdf
    import pymupdf
    doc1 = pymupdf.open(stream=r1.content, filetype="pdf")
    pages1 = len(doc1)
    doc1.close()
    assert pages1 == 2, f"1 record should produce 2 pages, got {pages1}"
    print(f"✅ PASS: 1 record → {pages1} pages (expected 2)")
    
    # Test 2 records → 2 pages
    payload2 = {
        "records": [
            payload1["records"][0],
            {
                "fio": "Петров Пётр Петрович",
                "birth_year": "2005",
                "stay_term": "срок до 30.06.2027",
                "from_place": "г. Гродно",
                "basis": "Договор найма № 6 от 22.07.2025",
                "area": "5467,9",
                "sign_date": "22.07.2025"
            }
        ]
    }
    
    r2 = requests.post(f"{BASE_URL}/zayavlenie/preview", json=payload2)
    assert r2.status_code == 200, f"Expected 200, got {r2.status_code}"
    doc2 = pymupdf.open(stream=r2.content, filetype="pdf")
    pages2 = len(doc2)
    doc2.close()
    assert pages2 == 2, f"2 records should produce 2 pages, got {pages2}"
    print(f"✅ PASS: 2 records → {pages2} pages (expected 2)")
    
    # Test 3 records → 4 pages
    payload3 = {
        "records": payload2["records"] + [{
            "fio": "Сидоров Сидор Сидорович",
            "birth_year": "2007",
            "stay_term": "срок до 30.06.2029",
            "from_place": "г. Брест",
            "basis": "Договор найма № 7 от 23.07.2025",
            "area": "5467,9",
            "sign_date": "23.07.2025"
        }]
    }
    
    r3 = requests.post(f"{BASE_URL}/zayavlenie/preview", json=payload3)
    assert r3.status_code == 200, f"Expected 200, got {r3.status_code}"
    doc3 = pymupdf.open(stream=r3.content, filetype="pdf")
    pages3 = len(doc3)
    doc3.close()
    assert pages3 == 4, f"3 records should produce 4 pages, got {pages3}"
    print(f"✅ PASS: 3 records → {pages3} pages (expected 4)")
    
    return True


def test_zayavlenie_preview_png():
    """TEST 3: POST /api/zayavlenie/preview-png → проверить side=front/back"""
    print("\n" + "="*80)
    print("TEST 3: POST /api/zayavlenie/preview-png (front/back)")
    print("="*80)
    
    payload = {
        "records": [{
            "fio": "Тестов Тест Тестович",
            "birth_year": "2006",
            "stay_term": "срок до 30.06.2028",
            "from_place": "г. Минск",
            "basis": "Договор найма № 1 от 01.01.2025",
            "area": "5467,9",
            "sign_date": "01.01.2025"
        }],
        "side": "front"
    }
    
    # Test front
    r_front = requests.post(f"{BASE_URL}/zayavlenie/preview-png", json=payload)
    assert r_front.status_code == 200, f"Expected 200, got {r_front.status_code}"
    assert r_front.headers.get("Content-Type") == "image/png", "Expected PNG"
    assert r_front.content[:4] == b'\x89PNG', "Response must be valid PNG"
    front_size = len(r_front.content)
    print(f"✅ PASS: side='front' → PNG {front_size} bytes")
    
    # Test back
    payload["side"] = "back"
    r_back = requests.post(f"{BASE_URL}/zayavlenie/preview-png", json=payload)
    assert r_back.status_code == 200, f"Expected 200, got {r_back.status_code}"
    assert r_back.headers.get("Content-Type") == "image/png", "Expected PNG"
    assert r_back.content[:4] == b'\x89PNG', "Response must be valid PNG"
    back_size = len(r_back.content)
    print(f"✅ PASS: side='back' → PNG {back_size} bytes")
    
    # Verify they are different
    assert r_front.content != r_back.content, "Front and back images must be different"
    print(f"✅ PASS: Front and back are DIFFERENT images (size diff: {abs(front_size - back_size)} bytes)")
    
    return True


def test_master_upload_flow():
    """TEST 4: GET /api/master-template → заполнить → POST /api/master-upload"""
    print("\n" + "="*80)
    print("TEST 4: Master template upload flow")
    print("="*80)
    
    # Download template
    r_tpl = requests.get(f"{BASE_URL}/master-template")
    assert r_tpl.status_code == 200, f"Expected 200, got {r_tpl.status_code}"
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in r_tpl.headers.get("Content-Type", "")
    print(f"✅ Downloaded template: {len(r_tpl.content)} bytes")
    
    # Load and fill template
    wb = load_workbook(io.BytesIO(r_tpl.content))
    ws = wb["Данные"] if "Данные" in wb.sheetnames else wb.active
    
    # Find header row (row 2 in the template)
    headers = {}
    for idx, cell in enumerate(ws[2], start=1):
        if cell.value:
            headers[str(cell.value).strip()] = idx
    
    print(f"   Found {len(headers)} headers")
    
    # Fill row 3 with test data
    row = 3
    test_data = {
        "ФИО (полностью)": "Иванов Иван Иванович",
        "Дата рождения (ДД.ММ.ГГГГ)": datetime(2006, 8, 3),  # datetime object
        "Паспорт (серия и номер)": "MP1234567",
        "Паспорт: кем выдан": "Первомайским РУВД г. Минска",
        "Паспорт: дата выдачи": datetime(2020, 1, 10),
        "Номер договора": "TEST-001",
        "Дата подписания договора": datetime(2025, 7, 21),
        "Срок договора до": datetime(2028, 6, 30),  # contract_end_date
        "Жительство: область": "Минская",
        "Жительство: район": "Минский",
        "Жительство: город (пгт)": "Минск",
        "Жительство: улица": "пр-т Дзержинского",
        "Жительство: дом": "85",
        "Жительство: комната/квартира": "302/2",
        "Откуда прибыл: область": "Гродненская",
        "Откуда прибыл: город (пгт)": "Лида",
    }
    
    for header, value in test_data.items():
        if header in headers:
            col_idx = headers[header]
            ws.cell(row=row, column=col_idx, value=value)
    
    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filled_xlsx = buf.getvalue()
    print(f"   Filled template: {len(filled_xlsx)} bytes")
    
    # Upload
    files = {"file": ("test_data.xlsx", filled_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r_upload = requests.post(f"{BASE_URL}/master-upload", files=files)
    assert r_upload.status_code == 200, f"Expected 200, got {r_upload.status_code}: {r_upload.text}"
    
    data = r_upload.json()
    assert "zayavlenie" in data, "Response must contain 'zayavlenie' key"
    assert isinstance(data["zayavlenie"], list), "'zayavlenie' must be a list"
    assert len(data["zayavlenie"]) > 0, "'zayavlenie' must not be empty"
    
    zayav = data["zayavlenie"][0]
    print(f"\n   Checking zayavlenie[0] fields:")
    
    # Check doc_name
    assert "doc_name" in zayav, "Missing 'doc_name'"
    assert zayav["doc_name"] == "паспорт гражданина Республики Беларусь", \
        f"doc_name should be 'паспорт гражданина Республики Беларусь', got '{zayav['doc_name']}'"
    print(f"   ✓ doc_name = '{zayav['doc_name']}'")
    
    # Check stay_term
    assert "stay_term" in zayav, "Missing 'stay_term'"
    assert zayav["stay_term"].startswith("срок до "), \
        f"stay_term should start with 'срок до ', got '{zayav['stay_term']}'"
    assert "00:00:00" not in zayav["stay_term"], \
        f"stay_term should NOT contain '00:00:00', got '{zayav['stay_term']}'"
    # Check date format ДД.ММ.ГГГГ
    import re
    date_match = re.search(r'\d{2}\.\d{2}\.\d{4}', zayav["stay_term"])
    assert date_match, f"stay_term should contain date in ДД.ММ.ГГГГ format, got '{zayav['stay_term']}'"
    print(f"   ✓ stay_term = '{zayav['stay_term']}'")
    
    # Check basis
    assert "basis" in zayav, "Missing 'basis'"
    assert zayav["basis"].startswith("Договор найма № "), \
        f"basis should start with 'Договор найма № ', got '{zayav['basis']}'"
    assert "TEST-001" in zayav["basis"], \
        f"basis should contain contract number 'TEST-001', got '{zayav['basis']}'"
    assert " от " in zayav["basis"], \
        f"basis should contain ' от ', got '{zayav['basis']}'"
    print(f"   ✓ basis = '{zayav['basis']}'")
    
    # Check birth_year
    assert "birth_year" in zayav, "Missing 'birth_year'"
    assert len(zayav["birth_year"]) == 4, \
        f"birth_year should be 4 digits, got '{zayav['birth_year']}'"
    assert zayav["birth_year"] == "2006", \
        f"birth_year should be '2006', got '{zayav['birth_year']}'"
    print(f"   ✓ birth_year = '{zayav['birth_year']}'")
    
    # Check area
    assert "area" in zayav, "Missing 'area'"
    assert zayav["area"] == "5467,9", \
        f"area should be '5467,9', got '{zayav['area']}'"
    print(f"   ✓ area = '{zayav['area']}'")
    
    # Check passport split
    assert "passport_series" in zayav, "Missing 'passport_series'"
    assert "passport_number" in zayav, "Missing 'passport_number'"
    assert zayav["passport_series"] == "MP", \
        f"passport_series should be 'MP', got '{zayav['passport_series']}'"
    assert zayav["passport_number"] == "1234567", \
        f"passport_number should be '1234567', got '{zayav['passport_number']}'"
    print(f"   ✓ passport_series = '{zayav['passport_series']}', passport_number = '{zayav['passport_number']}'")
    
    # Check passport_issue_date format
    assert "passport_issue_date" in zayav, "Missing 'passport_issue_date'"
    assert "00:00:00" not in zayav["passport_issue_date"], \
        f"passport_issue_date should NOT contain '00:00:00', got '{zayav['passport_issue_date']}'"
    print(f"   ✓ passport_issue_date = '{zayav['passport_issue_date']}'")
    
    print(f"\n✅ PASS: All zayavlenie[0] fields are correct")
    return data


def test_zayavlenie_prefill():
    """TEST 5: POST /api/zayavlenie/prefill"""
    print("\n" + "="*80)
    print("TEST 5: POST /api/zayavlenie/prefill")
    print("="*80)
    
    # Use master data from previous test
    master_data = {
        "fio": "Иванов Иван Иванович",
        "birth_date": "03.08.2006",
        "passport": "MP1234567",
        "passport_issued_by": "Первомайским РУВД г. Минска",
        "passport_issue_date": "10.01.2020",
        "contract_number": "TEST-001",
        "sign_date": "21.07.2025",
        "contract_end_date": "30.06.2028",
        "res_obl": "Минская",
        "res_raion": "Минский",
        "res_city": "Минск",
        "res_street": "пр-т Дзержинского",
        "res_house": "85",
        "res_apartment": "302/2",
        "from_obl": "Гродненская",
        "from_city": "Лида",
    }
    
    payload = {"students": [master_data]}
    r = requests.post(f"{BASE_URL}/zayavlenie/prefill", json=payload)
    assert r.status_code == 200, f"Expected 200, got {r.status_code}: {r.text}"
    
    data = r.json()
    assert "records" in data, "Response must contain 'records'"
    assert len(data["records"]) > 0, "'records' must not be empty"
    
    rec = data["records"][0]
    print(f"\n   Checking prefill record:")
    
    # Check same fields as master-upload
    assert rec.get("doc_name") == "паспорт гражданина Республики Беларусь"
    print(f"   ✓ doc_name = '{rec['doc_name']}'")
    
    assert rec.get("stay_term", "").startswith("срок до ")
    print(f"   ✓ stay_term = '{rec['stay_term']}'")
    
    assert rec.get("basis", "").startswith("Договор найма № ")
    print(f"   ✓ basis = '{rec['basis']}'")
    
    assert rec.get("birth_year") == "2006"
    print(f"   ✓ birth_year = '{rec['birth_year']}'")
    
    assert rec.get("area") == "5467,9"
    print(f"   ✓ area = '{rec['area']}'")
    
    assert rec.get("passport_series") == "MP"
    assert rec.get("passport_number") == "1234567"
    print(f"   ✓ passport split correctly")
    
    print(f"\n✅ PASS: Prefill produces correct fields")
    return True


def test_package_with_zayavlenie():
    """TEST 6: POST /api/package with zayavlenie"""
    print("\n" + "="*80)
    print("TEST 6: POST /api/package (with zayavlenie)")
    print("="*80)
    
    # First, upload master data to get proper master records
    print("   Uploading master data...")
    r_tpl = requests.get(f"{BASE_URL}/master-template")
    wb = load_workbook(io.BytesIO(r_tpl.content))
    ws = wb["Данные"] if "Данные" in wb.sheetnames else wb.active
    
    headers = {}
    for idx, cell in enumerate(ws[2], start=1):
        if cell.value:
            headers[str(cell.value).strip()] = idx
    
    row = 3
    test_data = {
        "ФИО (полностью)": "Тестов Тест Тестович",
        "Дата рождения (ДД.ММ.ГГГГ)": datetime(2006, 8, 3),
        "Паспорт (серия и номер)": "MP9999999",
        "Паспорт: кем выдан": "Тестовым РУВД",
        "Паспорт: дата выдачи": datetime(2020, 1, 10),
        "Номер договора": "PKG-001",
        "Дата подписания договора": datetime(2025, 7, 21),
        "Срок договора до": datetime(2028, 6, 30),
        "Жительство: город (пгт)": "Минск",
        "Жительство: улица": "ул. Тестовая",
        "Жительство: дом": "1",
        "Откуда прибыл: город (пгт)": "Гродно",
    }
    
    for header, value in test_data.items():
        if header in headers:
            ws.cell(row=row, column=headers[header], value=value)
    
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    files = {"file": ("pkg_test.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r_upload = requests.post(f"{BASE_URL}/master-upload", files=files)
    assert r_upload.status_code == 200, f"Upload failed: {r_upload.status_code}"
    
    master_data = r_upload.json()
    assert "master" in master_data, "Missing 'master' in upload response"
    
    # Test package with zayavlenie=true
    payload = {
        "people": master_data["master"],
        "duplex_flip": "long",
        "include": {
            "contract": False,
            "forma19": False,
            "forma24": False,
            "soobshenie": False,
            "zayavlenie": True
        }
    }
    
    r = requests.post(f"{BASE_URL}/package", json=payload)
    assert r.status_code == 200, f"Expected 200, got {r.status_code}: {r.text}"
    assert r.headers.get("Content-Type") == "application/pdf", "Expected PDF"
    assert r.content[:4] == b"%PDF", "Response must be valid PDF"
    
    # Check it's not empty
    assert len(r.content) > 1000, f"PDF too small: {len(r.content)} bytes"
    
    import pymupdf
    doc = pymupdf.open(stream=r.content, filetype="pdf")
    pages = len(doc)
    doc.close()
    
    print(f"✅ PASS: Package with zayavlenie → PDF {len(r.content)} bytes, {pages} pages")
    return True


def test_regression():
    """TEST 7: Regression tests"""
    print("\n" + "="*80)
    print("TEST 7: Regression tests")
    print("="*80)
    
    # Test root endpoint
    r1 = requests.get(f"{BASE_URL}/")
    assert r1.status_code == 200, f"GET / failed: {r1.status_code}"
    print("   ✓ GET /api/ → 200")
    
    # Test stats
    r2 = requests.get(f"{BASE_URL}/stats")
    assert r2.status_code == 200, f"GET /stats failed: {r2.status_code}"
    data = r2.json()
    assert "total" in data and "drafts" in data, "Stats missing required fields"
    print("   ✓ GET /api/stats → 200")
    
    # Test forma19/preview
    r3 = requests.post(f"{BASE_URL}/forma19/preview", json={
        "records": [{"surname": "Тест", "first_name": "Тестович"}],
        "duplex_flip": "long"
    })
    assert r3.status_code == 200, f"forma19/preview failed: {r3.status_code}"
    assert r3.content[:4] == b"%PDF", "forma19 must return PDF"
    print("   ✓ POST /api/forma19/preview → 200 PDF")
    
    # Test forma24/preview
    r4 = requests.post(f"{BASE_URL}/forma24/preview", json={
        "records": [{"surname": "Тест", "sex": "1"}],
        "duplex_flip": "long"
    })
    assert r4.status_code == 200, f"forma24/preview failed: {r4.status_code}"
    assert r4.content[:4] == b"%PDF", "forma24 must return PDF"
    print("   ✓ POST /api/forma24/preview → 200 PDF")
    
    # Test soobshenie (overlay/generate)
    r5 = requests.post(f"{BASE_URL}/overlay/generate", json={
        "records": [{"fio": "Тест"}],
        "page_size": "card",
        "with_form": True
    })
    assert r5.status_code == 200, f"overlay/generate failed: {r5.status_code}"
    assert r5.content[:4] == b"%PDF", "overlay must return PDF"
    print("   ✓ POST /api/overlay/generate → 200 PDF")
    
    # Test contracts/preview
    r6 = requests.post(f"{BASE_URL}/contracts/preview?format=pdf", json={
        "fields": {
            "contract_number": "REG-001",
            "full_name": "Регрессия Тест",
            "citizenship": "Республики Беларусь",
            "birth_date": "01.01.2000",
            "room_number": "101",
            "registration_address": "г. Минск",
            "passport_number": "AB1234567",
            "phone": "+375291234567"
        }
    })
    assert r6.status_code == 200, f"contracts/preview failed: {r6.status_code}"
    assert r6.content[:4] == b"%PDF", "contract must return PDF"
    print("   ✓ POST /api/contracts/preview → 200 PDF")
    
    print(f"\n✅ PASS: All regression tests passed")
    return True


def main():
    print("\n" + "="*80)
    print("ТЕСТИРОВАНИЕ: Заявление о регистрации по месту ПРЕБЫВАНИЯ")
    print("Backend: " + BASE_URL)
    print("="*80)
    
    tests = [
        ("GET /api/zayavlenie/fields", test_zayavlenie_fields),
        ("POST /api/zayavlenie/preview (pages)", test_zayavlenie_preview_pages),
        ("POST /api/zayavlenie/preview-png", test_zayavlenie_preview_png),
        ("Master template upload flow", test_master_upload_flow),
        ("POST /api/zayavlenie/prefill", test_zayavlenie_prefill),
        ("POST /api/package", test_package_with_zayavlenie),
        ("Regression tests", test_regression),
    ]
    
    passed = 0
    failed = 0
    
    for name, test_func in tests:
        try:
            test_func()
            passed += 1
        except AssertionError as e:
            print(f"\n❌ FAIL: {name}")
            print(f"   Error: {e}")
            failed += 1
        except Exception as e:
            print(f"\n❌ ERROR: {name}")
            print(f"   Exception: {e}")
            import traceback
            traceback.print_exc()
            failed += 1
    
    print("\n" + "="*80)
    print(f"SUMMARY: {passed} passed, {failed} failed out of {len(tests)} tests")
    print("="*80)
    
    if failed > 0:
        sys.exit(1)
    else:
        print("\n✅ ALL TESTS PASSED")
        sys.exit(0)


if __name__ == "__main__":
    main()
