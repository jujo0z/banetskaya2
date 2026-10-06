#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""СВЕРКА (только чтение) файла заселения с главной базой contracts.
Сравнивает номера договоров, сроки, комнату по ФИО."""
import re
from difflib import SequenceMatcher
from datetime import datetime, date
from openpyxl import load_workbook
from pymongo import MongoClient

XLSX="/root/dlya_ivana.xlsx"
db=MongoClient("mongodb://127.0.0.1:27017")["banetskaya_db"]

def _clean(v):
    if v is None: return ""
    if isinstance(v,(datetime,date)): return v.strftime("%d.%m.%Y")
    s=str(v).strip()
    m=re.match(r"^(\d{4})-(\d{2})-(\d{2})(?:\s+00:00:00)?$",s)
    if m: return f"{m.group(3)}.{m.group(2)}.{m.group(1)}"
    return s
def norm(f):
    return re.sub(r"\s+"," ",(f or "").lower().replace("ё","е").strip())
def digits(s):
    return re.sub(r"\D","",s or "")

wb=load_workbook(XLSX,data_only=True); ws=wb[wb.sheetnames[0]]
rows=list(ws.iter_rows(values_only=True))
people=[]
for r in rows[3:]:
    def g(i): return _clean(r[i]) if i<len(r) else ""
    if not g(2): continue
    block=g(1)
    people.append({"fio":g(2),"block":block,"group":g(4),"num":g(6),"term":g(7)})

docs=list(db.contracts.find({},{"_id":0}))
by_norm={}
for d in docs:
    by_norm.setdefault(norm(d.get("full_name","")),[]).append(d)

def toks(n): return norm(n).split()
def fuzzy(fio, free):
    pt=toks(fio); ps=pt[0] if pt else ""; pf=pt[1] if len(pt)>1 else ""
    best=None; bs=0
    for d in free:
        dt=toks(d.get("full_name","")); ds=dt[0] if dt else ""; df=dt[1] if len(dt)>1 else ""
        fr=SequenceMatcher(None,norm(fio),norm(d.get("full_name",""))).ratio()
        sr=SequenceMatcher(None,ps,ds).ratio()
        sc=0
        if pf and pf==df and sr>=0.6: sc=1.5+sr
        elif fr>=0.9: sc=fr
        if sc>bs: best,bs=d,sc
    return best

used=set(); matched=0
num_diff=[]; term_diff=[]; room_diff=[]; not_in_base=[]
for p in people:
    cands=[d for d in by_norm.get(norm(p["fio"]),[]) if id(d) not in used]
    d=cands[0] if cands else fuzzy(p["fio"],[x for x in docs if id(x) not in used])
    if not d:
        not_in_base.append(p); continue
    used.add(id(d)); matched+=1
    m=d.get("master") or {}
    base_num=d.get("contract_number") or m.get("contract_number","")
    base_term=m.get("contract_end_date","")
    base_room=m.get("room_number","") or m.get("res_apartment","")
    if digits(p["num"])!=digits(base_num):
        num_diff.append((p["fio"],base_num,p["num"]))
    if p["term"] and base_term and p["term"]!=base_term:
        term_diff.append((p["fio"],base_term,p["term"]))
    if p["block"] and base_room and p["block"].replace(" ","")!=base_room.replace(" ",""):
        room_diff.append((p["fio"],base_room,p["block"]))
not_in_file=[d for d in docs if id(d) not in used]

print("="*60)
print("СВЕРКА файла с главной базой contracts (только чтение)")
print("="*60)
print("Студентов в файле:           ", len(people))
print("Записей в базе contracts:    ", len(docs))
print("Сопоставлено по ФИО:         ", matched)
print("Из файла НЕТ в базе:         ", len(not_in_base))
print("В базе есть, но НЕТ в файле:  ", len(not_in_file))
print("-"*60)
print("Расходится НОМЕР договора:   ", len(num_diff))
print("Расходится СРОК:             ", len(term_diff))
print("Расходится КОМНАТА:          ", len(room_diff))
print("="*60)
print("\n--- ИЗ ФАЙЛА НЕТ В БАЗЕ contracts [первые 50] ---")
for p in not_in_base[:50]: print(f'  {p["block"]:8} {p["fio"]}  дог:{p["num"]}')
if len(not_in_base)>50: print("  ... ещё", len(not_in_base)-50)
print("\n--- РАСХОЖДЕНИЕ НОМЕРА договора (база -> файл) [первые 40] ---")
for x in num_diff[:40]: print(f'  {x[0]}: {x[1]!r} -> {x[2]!r}')
if len(num_diff)>40: print("  ... ещё", len(num_diff)-40)
print("\n--- РАСХОЖДЕНИЕ СРОКА (база -> файл) [первые 40] ---")
for x in term_diff[:40]: print(f'  {x[0]}: {x[1]} -> {x[2]}')
if len(term_diff)>40: print("  ... ещё", len(term_diff)-40)
print("\n--- В БАЗЕ, НО НЕТ В ФАЙЛЕ [первые 40] ---")
for d in not_in_file[:40]: print(f'  {d.get("full_name")}  дог:{d.get("contract_number")}')
if len(not_in_file)>40: print("  ... ещё", len(not_in_file)-40)
