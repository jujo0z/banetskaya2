#!/usr/bin/env python3
"""
Backend API testing for new master-template, master-upload, and package endpoints.
Tests the unified Excel template and full document package generation.
"""
import os
import sys
import requests
import json
import io
from openpyxl import load_workbook

# Get backend URL from environment
BACKEND_URL = os.environ.get('REACT_APP_BACKEND_URL', 'http://localhost:8001')
API_BASE = f"{BACKEND_URL}/api"

print(f"Testing backend at: {API_BASE}")

def test_master_template():
    """Test 1: GET /api/master-template returns valid Excel with correct structure"""
    print("\n[TEST 1] GET /api/master-template")
    url = f"{API_BASE}/master-template"
    
    try:
        resp = requests.get(url, timeout=30)
        print(f"  Status: {resp.status_code}")
        print(f"  Content-Type: {resp.headers.get('Content-Type')}")
        
        assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"
        assert 'spreadsheet' in resp.headers.get('Content-Type', ''), f"Expected Excel content type, got {resp.headers.get('Content-Type')}"
        
        # Load Excel with openpyxl
        excel_bytes = resp.content
        print(f"  Excel size: {len(excel_bytes)} bytes")
        
        wb = load_workbook(io.BytesIO(excel_bytes))
        print(f"  Sheets: {wb.sheetnames}")
        
        # Check required sheets
        assert "Данные" in wb.sheetnames, "Sheet 'Данные' not found"
        assert "Инструкция" in wb.sheetnames, "Sheet 'Инструкция' not found"
        
        # Check "Данные" sheet structure
        ws = wb["Данные"]
        
        # Row 2 should have headers (e.g., "ФИО (полностью)")
        row2_values = [cell.value for cell in ws[2]]
        print(f"  Row 2 (headers) first 3 cells: {row2_values[:3]}")
        assert any("ФИО" in str(v) for v in row2_values if v), "Header 'ФИО (полностью)' not found in row 2"
        
        # Row 3 should have sample data (e.g., "Иванов Иван Иванович")
        row3_values = [cell.value for cell in ws[3]]
        print(f"  Row 3 (sample) first 3 cells: {row3_values[:3]}")
        assert any("Иванов" in str(v) for v in row3_values if v), "Sample data 'Иванов Иван Иванович' not found in row 3"
        
        print("  ✅ PASS: Excel structure is correct")
        return excel_bytes
    except Exception as e:
        print(f"  ❌ FAIL: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


def test_master_upload(excel_bytes):
    """Test 2: POST /api/master-upload with the downloaded Excel returns parsed data"""
    print("\n[TEST 2] POST /api/master-upload")
    url = f"{API_BASE}/master-upload"
    
    try:
        files = {'file': ('test.xlsx', excel_bytes, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        resp = requests.post(url, files=files, timeout=30)
        print(f"  Status: {resp.status_code}")
        print(f"  Content-Type: {resp.headers.get('Content-Type')}")
        
        assert resp.status_code == 200, f"Expected 200, got {resp.status_code}: {resp.text}"
        
        data = resp.json()
        print(f"  Response keys: {list(data.keys())}")
        
        # Check required keys
        assert "count" in data, "Key 'count' not found"
        assert "master" in data, "Key 'master' not found"
        assert "contracts" in data, "Key 'contracts' not found"
        assert "forma19" in data, "Key 'forma19' not found"
        assert "forma24" in data, "Key 'forma24' not found"
        assert "soobshenie" in data, "Key 'soobshenie' not found"
        
        print(f"  count: {data['count']}")
        assert data['count'] >= 1, f"Expected count >= 1, got {data['count']}"
        
        # Check arrays are non-empty
        assert len(data['master']) >= 1, "master array is empty"
        assert len(data['contracts']) >= 1, "contracts array is empty"
        assert len(data['forma19']) >= 1, "forma19 array is empty"
        assert len(data['forma24']) >= 1, "forma24 array is empty"
        assert len(data['soobshenie']) >= 1, "soobshenie array is empty"
        
        print(f"  master[0] keys: {list(data['master'][0].keys())[:5]}...")
        
        # Check forma24[0] structure
        f24 = data['forma24'][0]
        print(f"  forma24[0] keys: {list(f24.keys())[:10]}...")
        assert "surname" in f24, "forma24[0] missing 'surname'"
        assert f24["surname"] == "Иванов", f"Expected surname='Иванов', got {f24['surname']}"
        assert "first_name" in f24, "forma24[0] missing 'first_name'"
        assert f24["first_name"] == "Иван", f"Expected first_name='Иван', got {f24['first_name']}"
        assert "birth_month" in f24, "forma24[0] missing 'birth_month'"
        assert f24["birth_month"] in ["марта", "март"], f"Expected birth_month='марта', got {f24['birth_month']}"
        assert "sex" in f24, "forma24[0] missing 'sex'"
        assert f24["sex"] == "1", f"Expected sex='1', got {f24['sex']}"
        
        # Check contracts[0] structure
        contract = data['contracts'][0]
        print(f"  contracts[0] keys: {list(contract.keys())[:10]}...")
        assert "full_name" in contract, "contracts[0] missing 'full_name'"
        assert contract["full_name"] == "Иванов Иван Иванович", f"Expected full_name='Иванов Иван Иванович', got {contract['full_name']}"
        assert "registration_address" in contract, "contracts[0] missing 'registration_address'"
        assert len(contract["registration_address"]) > 0, "registration_address is empty"
        print(f"  contracts[0] registration_address: {contract['registration_address'][:50]}...")
        
        # Check soobshenie[0] structure
        soob = data['soobshenie'][0]
        print(f"  soobshenie[0] keys: {list(soob.keys())[:10]}...")
        assert "fio" in soob, "soobshenie[0] missing 'fio'"
        assert "date_day" in soob, "soobshenie[0] missing 'date_day'"
        assert "passport_series" in soob, "soobshenie[0] missing 'passport_series'"
        assert soob["passport_series"] == "MP", f"Expected passport_series='MP', got {soob['passport_series']}"
        
        print("  ✅ PASS: Upload and parsing successful")
        return data
    except Exception as e:
        print(f"  ❌ FAIL: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


def test_package_full(upload_data):
    """Test 3: POST /api/package with full document set returns valid PDF with 7 pages per person"""
    print("\n[TEST 3] POST /api/package (full document set)")
    url = f"{API_BASE}/package"
    
    try:
        # Use only first person from master data
        payload = {
            "people": [upload_data['master'][0]],
            "duplex_flip": "long"
        }
        
        resp = requests.post(url, json=payload, timeout=60)
        print(f"  Status: {resp.status_code}")
        print(f"  Content-Type: {resp.headers.get('Content-Type')}")
        
        assert resp.status_code == 200, f"Expected 200, got {resp.status_code}: {resp.text}"
        assert resp.headers.get('Content-Type') == 'application/pdf', f"Expected application/pdf, got {resp.headers.get('Content-Type')}"
        
        pdf_bytes = resp.content
        print(f"  PDF size: {len(pdf_bytes)} bytes")
        assert pdf_bytes[:4] == b'%PDF', f"Expected PDF signature, got {pdf_bytes[:4]}"
        
        # Check PDF structure with pymupdf
        import pymupdf
        doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
        page_count = len(doc)
        print(f"  Page count: {page_count}")
        
        # Expected: 1 person = 7 pages
        # - Contract: 4 pages
        # - Combined Forma19+Forma24: 2 pages (front + back)
        # - Soobshenie: 1 page
        assert page_count == 7, f"Expected 7 pages for 1 person, got {page_count}"
        
        # Check page 0 contains "ДОГОВОР"
        page0_text = doc[0].get_text()
        assert "ДОГОВОР" in page0_text, "Page 0 should contain 'ДОГОВОР'"
        print(f"  Page 0 contains: 'ДОГОВОР' ✓")
        
        # Check page 4 contains BOTH "Форма № 19" AND "Форма 24"
        page4_text = doc[4].get_text()
        has_f19 = "Форма № 19" in page4_text or "Форма №19" in page4_text or "АДРЕСНЫЙ ЛИСТОК" in page4_text
        has_f24 = "Форма 24" in page4_text or "Форма №24" in page4_text or "ТАЛОН МИГРАЦИОННОГО" in page4_text
        print(f"  Page 4 text sample: {page4_text[:200]}...")
        assert has_f19, "Page 4 should contain 'Форма № 19' or 'АДРЕСНЫЙ ЛИСТОК'"
        assert has_f24, "Page 4 should contain 'Форма 24' or 'ТАЛОН МИГРАЦИОННОГО'"
        print(f"  Page 4 contains: 'Форма № 19' ✓ and 'Форма 24' ✓")
        
        # Check page 6 contains "СООБЩЕНИЕ"
        page6_text = doc[6].get_text()
        assert "СООБЩЕНИЕ" in page6_text, "Page 6 should contain 'СООБЩЕНИЕ'"
        print(f"  Page 6 contains: 'СООБЩЕНИЕ' ✓")
        
        doc.close()
        print("  ✅ PASS: Full package PDF is correct")
        return pdf_bytes
    except Exception as e:
        print(f"  ❌ FAIL: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


def test_package_partial(upload_data):
    """Test 4: POST /api/package with partial document set (only forma19+forma24) returns 2 pages"""
    print("\n[TEST 4] POST /api/package (partial: only forma19+forma24)")
    url = f"{API_BASE}/package"
    
    try:
        payload = {
            "people": [upload_data['master'][0]],
            "duplex_flip": "long",
            "include": {
                "contract": False,
                "forma19": True,
                "forma24": True,
                "soobshenie": False
            }
        }
        
        resp = requests.post(url, json=payload, timeout=60)
        print(f"  Status: {resp.status_code}")
        print(f"  Content-Type: {resp.headers.get('Content-Type')}")
        
        assert resp.status_code == 200, f"Expected 200, got {resp.status_code}: {resp.text}"
        assert resp.headers.get('Content-Type') == 'application/pdf', f"Expected application/pdf, got {resp.headers.get('Content-Type')}"
        
        pdf_bytes = resp.content
        print(f"  PDF size: {len(pdf_bytes)} bytes")
        assert pdf_bytes[:4] == b'%PDF', f"Expected PDF signature, got {pdf_bytes[:4]}"
        
        # Check PDF structure with pymupdf
        import pymupdf
        doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
        page_count = len(doc)
        print(f"  Page count: {page_count}")
        
        # Expected: only combined forma19+forma24 = 2 pages (front + back)
        assert page_count == 2, f"Expected 2 pages (combined forma19+forma24), got {page_count}"
        
        doc.close()
        print("  ✅ PASS: Partial package PDF is correct")
    except Exception as e:
        print(f"  ❌ FAIL: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


def test_regression_fields():
    """Test 5: GET /api/fields (regression test)"""
    print("\n[TEST 5] GET /api/fields (regression)")
    url = f"{API_BASE}/fields"
    
    try:
        resp = requests.get(url, timeout=30)
        print(f"  Status: {resp.status_code}")
        
        assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"
        
        data = resp.json()
        assert "fields" in data, "Expected 'fields' key in response"
        assert isinstance(data["fields"], list), "Expected 'fields' to be a list"
        assert len(data["fields"]) > 0, "Fields list is empty"
        print(f"  Fields count: {len(data['fields'])}")
        print("  ✅ PASS: Fields endpoint working")
    except Exception as e:
        print(f"  ❌ FAIL: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


def test_regression_forma24_preview():
    """Test 6: POST /api/forma24/preview (regression test)"""
    print("\n[TEST 6] POST /api/forma24/preview (regression)")
    url = f"{API_BASE}/forma24/preview"
    
    try:
        payload = {
            "records": [{"surname": "Тест"}],
            "duplex_flip": "long"
        }
        
        resp = requests.post(url, json=payload, timeout=30)
        print(f"  Status: {resp.status_code}")
        print(f"  Content-Type: {resp.headers.get('Content-Type')}")
        
        assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"
        assert resp.headers.get('Content-Type') == 'application/pdf', f"Expected application/pdf, got {resp.headers.get('Content-Type')}"
        
        pdf_bytes = resp.content
        assert pdf_bytes[:4] == b'%PDF', f"Expected PDF signature, got {pdf_bytes[:4]}"
        print(f"  PDF size: {len(pdf_bytes)} bytes")
        print("  ✅ PASS: Forma24 preview working")
    except Exception as e:
        print(f"  ❌ FAIL: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    print("=" * 80)
    print("BACKEND TESTING: Master Template & Package Endpoints")
    print("=" * 80)
    
    # Run tests in sequence
    excel_bytes = test_master_template()
    upload_data = test_master_upload(excel_bytes)
    test_package_full(upload_data)
    test_package_partial(upload_data)
    test_regression_fields()
    test_regression_forma24_preview()
    
    print("\n" + "=" * 80)
    print("✅ ALL TESTS PASSED (6/6)")
    print("=" * 80)
    sys.exit(0)
