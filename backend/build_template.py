"""Generates a default договор найма .docx template with docxtpl (Jinja) placeholders.
Run: python build_template.py  -> creates templates/contract_template.docx
The admin can later replace this file with the real college template (keeping the {{ }} placeholders).
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE_DIR = Path(__file__).parent / "templates"
TEMPLATE_DIR.mkdir(exist_ok=True)
OUT = TEMPLATE_DIR / "contract_template.docx"


def set_base_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)


def p(doc, text, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, size=12, space_after=6):
    para = doc.add_paragraph()
    para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    return para


def build():
    doc = Document()
    set_base_style(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    p(doc, "ДОГОВОР НАЙМА ЖИЛОГО ПОМЕЩЕНИЯ В ОБЩЕЖИТИИ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, space_after=2)
    p(doc, "№ {{ contract_number }}", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=10)

    header = doc.add_paragraph()
    header.paragraph_format.space_after = Pt(12)
    left = header.add_run("г. Алматы")
    left.font.name = "Times New Roman"
    left.font.size = Pt(12)
    header.add_run("\t\t\t\t\t\t")
    right = header.add_run("«{{ sign_date }}»")
    right.font.name = "Times New Roman"
    right.font.size = Pt(12)

    p(doc,
      "Колледж «Alatau Business College», именуемый в дальнейшем «Наймодатель», "
      "в лице директора Абдиева А.А., действующего на основании Устава, с одной стороны, "
      "и гражданин(-ка) {{ citizenship }} {{ full_name }}, {{ birth_date }} года рождения, "
      "именуемый(-ая) в дальнейшем «Наниматель», с другой стороны, на основании приказа "
      "№ {{ order_number }} от {{ order_date }}, заключили настоящий договор о нижеследующем:")

    p(doc, "1. ПРЕДМЕТ ДОГОВОРА", bold=True, space_after=4)
    p(doc,
      "1.1. Наймодатель предоставляет Нанимателю за плату во временное пользование "
      "койко-место в жилом помещении (комната № {{ room_number }}, площадь 18 кв.м) "
      "в общежитии колледжа, расположенном по адресу: г. Алматы, ул. Толе би, 187.")
    p(doc,
      "1.2. Жилое помещение предоставляется для проживания на период обучения "
      "сроком до {{ contract_end_date }}.")

    p(doc, "2. ДАННЫЕ НАНИМАТЕЛЯ", bold=True, space_after=4)
    p(doc, "2.1. Ф.И.О.: {{ full_name }}")
    p(doc, "2.2. Дата рождения: {{ birth_date }}")
    p(doc, "2.3. Гражданство: {{ citizenship }}")
    p(doc, "2.4. Адрес регистрации: {{ registration_address }}")
    p(doc, "2.5. Документ, удостоверяющий личность: № {{ passport_number }}, "
           "выдан {{ passport_issued_by }} {{ passport_issue_date }}, "
           "срок действия до {{ passport_valid_until }}.")
    p(doc, "2.6. Идентификационный номер (ИИН): {{ id_number }}")
    p(doc, "2.7. Контактный телефон: {{ phone }}")

    p(doc, "3. ПРАВА И ОБЯЗАННОСТИ СТОРОН", bold=True, space_after=4)
    p(doc, "3.1. Наниматель обязуется использовать жилое помещение по назначению, "
           "соблюдать правила внутреннего распорядка общежития, бережно относиться к "
           "имуществу и своевременно вносить плату за проживание.")
    p(doc, "3.2. Наймодатель обязуется обеспечить надлежащее санитарно-техническое "
           "состояние жилого помещения и предоставить коммунальные услуги.")

    p(doc, "4. СРОК ДЕЙСТВИЯ ДОГОВОРА", bold=True, space_after=4)
    p(doc, "4.1. Настоящий договор вступает в силу с момента подписания и действует "
           "до {{ contract_end_date }}.")

    p(doc, "5. РЕКВИЗИТЫ И ПОДПИСИ СТОРОН", bold=True, space_after=8)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    c1, c2 = table.rows[0].cells

    def cell_lines(cell, lines):
        cell.paragraphs[0].text = ""
        for i, ln in enumerate(lines):
            par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            run = par.add_run(ln)
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            if ln.startswith("НАЙМ") or ln.startswith("НАН"):
                run.bold = True

    cell_lines(c1, [
        "НАЙМОДАТЕЛЬ:",
        "Колледж «Alatau Business College»",
        "г. Алматы, ул. Толе би, 187",
        "БИН 123456789012",
        "",
        "Директор _______________ Абдиев А.А.",
        "",
        "Начальник ОКЮР ________ Сериков С.С.",
    ])
    cell_lines(c2, [
        "НАНИМАТЕЛЬ:",
        "{{ full_name }}",
        "{{ registration_address }}",
        "ИИН {{ id_number }}",
        "Тел. {{ phone }}",
        "",
        "_______________ / подпись /",
    ])

    doc.save(OUT)
    print(f"Template saved to {OUT}")


if __name__ == "__main__":
    build()
