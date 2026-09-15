"""End-to-end backend tests for contract generator API."""
import io
import os
import zipfile
import pytest
import requests
from openpyxl import Workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://deploy-preview-152.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

RUSSIAN_HEADERS = [
    "Номер договора", "Дата подписания", "Номер приказа", "Дата приказа",
    "Гражданство", "ФИО", "Дата рождения", "Номер комнаты",
    "Срок договора до", "Адрес регистрации", "Номер паспорта", "Дата выдачи",
    "Срок действия", "Кем выдан", "ИИН", "Телефон",
]

SAMPLE_ROWS = [
    ["1001", "01.09.2025", "П-100", "20.08.2025", "РК", "Иванов Иван Иванович",
     "01.01.2005", "301", "30.06.2026", "г. Алматы, ул. Абая 1", "N12345678",
     "01.01.2020", "01.01.2030", "МВД РК", "050101300123", "+77001234567"],
    ["1002", "02.09.2025", "П-101", "20.08.2025", "РК", "Петров Пётр Петрович",
     "05.05.2005", "302", "30.06.2026", "г. Астана, ул. Мира 5", "N87654321",
     "01.02.2020", "01.02.2030", "МВД РК", "050505300234", "+77007654321"],
]


def _make_xlsx_bytes(headers=RUSSIAN_HEADERS, rows=SAMPLE_ROWS):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


@pytest.fixture(scope="module")
def created_ids():
    return []


# ---------- Basic ----------
def test_root():
    r = requests.get(f"{API}/")
    assert r.status_code == 200
    assert "message" in r.json()


def test_get_fields():
    r = requests.get(f"{API}/fields")
    assert r.status_code == 200
    data = r.json()
    assert "fields" in data
    assert len(data["fields"]) == 17
    keys = {f["key"] for f in data["fields"]}
    assert "full_name" in keys and "contract_number" in keys


# ---------- Upload ----------
def test_upload_xlsx_and_automap_all_16():
    xlsx = _make_xlsx_bytes()
    files = {"file": ("students.xlsx", xlsx,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{API}/upload", files=files)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["columns"] == RUSSIAN_HEADERS
    assert len(data["rows"]) == 2
    mapping = data["mapping"]
    # All 16 fields should be mapped
    assert len(mapping) == 16, f"Only mapped {len(mapping)}: {mapping}"
    # Spot checks
    assert mapping["full_name"] == "ФИО"
    assert mapping["contract_number"] == "Номер договора"
    assert mapping["phone"] == "Телефон"
    assert mapping["id_number"] == "ИИН"


def test_upload_non_xlsx_returns_400():
    files = {"file": ("bad.txt", b"not an excel", "text/plain")}
    r = requests.post(f"{API}/upload", files=files)
    assert r.status_code == 400


# ---------- Contracts CRUD ----------
def _sample_fields(num="TEST_9001", name="TEST_Иванов Иван"):
    return {
        "contract_number": num, "sign_date": "01.09.2025", "order_number": "П-1",
        "order_date": "20.08.2025", "citizenship": "РК", "full_name": name,
        "birth_date": "01.01.2005", "room_number": "301", "contract_end_date": "30.06.2026",
        "registration_address": "г. Алматы", "passport_number": "N123", "passport_issue_date": "01.01.2020",
        "passport_valid_until": "01.01.2030", "passport_issued_by": "МВД", "id_number": "050101300123",
        "phone": "+77001234567",
    }


def test_create_contract(created_ids):
    r = requests.post(f"{API}/contracts", json={"fields": _sample_fields()})
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["id"] and data["contract_number"] == "TEST_9001"
    assert data["full_name"] == "TEST_Иванов Иван"
    created_ids.append(data["id"])


def test_batch_create(created_ids):
    contracts = [_sample_fields(num=f"TEST_9{i:03d}", name=f"TEST_Студент {i}") for i in range(2, 5)]
    r = requests.post(f"{API}/contracts/batch", json={"contracts": contracts})
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["count"] == 3
    for c in data["created"]:
        created_ids.append(c["id"])


def test_list_and_search(created_ids):
    r = requests.get(f"{API}/contracts")
    assert r.status_code == 200
    assert isinstance(r.json(), list)

    r = requests.get(f"{API}/contracts", params={"q": "TEST_9001"})
    assert r.status_code == 200
    results = r.json()
    assert any(c["contract_number"] == "TEST_9001" for c in results)

    # case-insensitive full_name
    r = requests.get(f"{API}/contracts", params={"q": "test_иванов"})
    assert r.status_code == 200
    assert len(r.json()) >= 1


def test_download_docx(created_ids):
    cid = created_ids[0]
    r = requests.get(f"{API}/contracts/{cid}/download", params={"format": "docx"})
    assert r.status_code == 200, r.text
    assert "wordprocessingml" in r.headers.get("content-type", "")
    content = r.content
    assert content[:2] == b"PK", "Not a valid docx (zip)"
    # verify no leftover {{ }} jinja placeholders in document.xml
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        with zf.open("word/document.xml") as f:
            xml = f.read().decode("utf-8", errors="ignore")
    assert "{{" not in xml and "}}" not in xml, "Unrendered placeholders present"
    assert "TEST_Иванов Иван" in xml or "TEST_9001" in xml


def test_download_pdf(created_ids):
    cid = created_ids[0]
    r = requests.get(f"{API}/contracts/{cid}/download", params={"format": "pdf"}, timeout=120)
    assert r.status_code == 200, r.text
    assert r.headers.get("content-type", "").startswith("application/pdf")
    assert r.content[:4] == b"%PDF"


def test_preview_docx():
    r = requests.post(f"{API}/contracts/preview", params={"format": "docx"},
                      json={"fields": _sample_fields(num="TEST_PREV", name="TEST_Preview")})
    assert r.status_code == 200
    assert r.content[:2] == b"PK"


def test_batch_download_zip(created_ids):
    ids = created_ids[:3]
    r = requests.post(f"{API}/contracts/batch-download",
                      json={"ids": ids, "format": "docx"}, timeout=120)
    assert r.status_code == 200, r.text
    assert r.headers.get("content-type", "").startswith("application/zip")
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        names = zf.namelist()
    assert len(names) == len(ids)
    assert all(n.endswith(".docx") for n in names)


def test_batch_download_pdf_zip(created_ids):
    ids = created_ids[:2]
    r = requests.post(f"{API}/contracts/batch-download",
                      json={"ids": ids, "format": "pdf"}, timeout=180)
    assert r.status_code == 200, r.text
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        names = zf.namelist()
    assert len(names) == 2 and all(n.endswith(".pdf") for n in names)


def test_delete_contract(created_ids):
    # delete all created
    for cid in list(created_ids):
        r = requests.delete(f"{API}/contracts/{cid}")
        assert r.status_code == 200
    # verify one is 404
    r = requests.get(f"{API}/contracts/{created_ids[0]}")
    assert r.status_code == 404
    created_ids.clear()


# ---------- New features: stats, note field, batch-print, sample-template ----------
def test_stats_endpoint():
    r = requests.get(f"{API}/stats")
    assert r.status_code == 200
    data = r.json()
    for k in ("total", "this_month", "datasets", "recent"):
        assert k in data
    assert isinstance(data["recent"], list)
    assert isinstance(data["total"], int)


def test_sample_template_download():
    r = requests.get(f"{API}/sample-template")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("content-type", "")
    assert r.content[:2] == b"PK"


def test_note_field_appears_when_filled():
    fields = _sample_fields(num="TEST_NOTE1", name="TEST_NoteUser")
    fields["note"] = "TEST_ManualNoteXYZ"
    r = requests.post(f"{API}/contracts/preview", params={"format": "docx"}, json={"fields": fields})
    assert r.status_code == 200
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    assert "TEST_ManualNoteXYZ" in xml
    assert "Дополнительно" in xml


def test_note_field_absent_when_empty():
    fields = _sample_fields(num="TEST_NOTE2", name="TEST_NoNote")
    # explicitly empty
    fields["note"] = ""
    r = requests.post(f"{API}/contracts/preview", params={"format": "docx"}, json={"fields": fields})
    assert r.status_code == 200
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    # 'Дополнительно:' line must not appear
    assert "Дополнительно:" not in xml
    # no leftover placeholders
    assert "{{" not in xml and "}}" not in xml


def test_batch_print_merged_pdf():
    # create 2 contracts
    ids = []
    for i in range(2):
        r = requests.post(f"{API}/contracts", json={"fields": _sample_fields(num=f"TEST_MP{i}", name=f"TEST_Merge{i}")})
        assert r.status_code == 200
        ids.append(r.json()["id"])
    try:
        r = requests.post(f"{API}/contracts/batch-print", json={"ids": ids, "format": "pdf"}, timeout=180)
        assert r.status_code == 200, r.text
        assert r.headers.get("content-type", "").startswith("application/pdf")
        assert r.content[:4] == b"%PDF"
        # verify it has at least 2 pages
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(r.content))
        assert len(reader.pages) >= 2
    finally:
        for cid in ids:
            requests.delete(f"{API}/contracts/{cid}")


# ---------- Moderation (extra_*) ----------
EXTRA_KEYS = ["extra_subject", "extra_tenant", "extra_landlord",
              "extra_liability", "extra_term", "extra_other"]


def test_moderation_extra_fields_appear_when_filled():
    fields = _sample_fields(num="TEST_MOD1", name="TEST_ModUser")
    markers = {k: f"TEST_MARKER_{k.upper()}" for k in EXTRA_KEYS}
    fields.update(markers)
    r = requests.post(f"{API}/contracts/preview", params={"format": "docx"}, json={"fields": fields})
    assert r.status_code == 200
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    for k, marker in markers.items():
        assert marker in xml, f"Marker for {k} not in document: {marker}"
    # No leftover jinja/docxtpl syntax
    assert "{{" not in xml and "}}" not in xml
    assert "{%p" not in xml and "%}" not in xml


def test_moderation_empty_extras_do_not_appear():
    fields = _sample_fields(num="TEST_MOD2", name="TEST_NoMod")
    # explicit empty extras
    for k in EXTRA_KEYS:
        fields[k] = ""
    r = requests.post(f"{API}/contracts/preview", params={"format": "docx"}, json={"fields": fields})
    assert r.status_code == 200
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    assert "TEST_MARKER" not in xml
    assert "{{" not in xml and "}}" not in xml
    assert "{%p" not in xml and "%}" not in xml
    # No 'Дополнительно' block since note also empty
    assert "Дополнительно:" not in xml


def test_moderation_persists_via_history():
    """POST /api/contracts with extra_* saves them, and downloading later reproduces them."""
    fields = _sample_fields(num="TEST_MOD_H", name="TEST_ModHist")
    fields["extra_subject"] = "TEST_HIST_SUBJECT_XYZ"
    fields["extra_other"] = "TEST_HIST_OTHER_ABC"
    fields["note"] = "TEST_HIST_NOTE_QQQ"
    r = requests.post(f"{API}/contracts", json={"fields": fields})
    assert r.status_code == 200
    cid = r.json()["id"]
    try:
        # Verify stored fields via GET
        r2 = requests.get(f"{API}/contracts/{cid}")
        assert r2.status_code == 200
        stored = r2.json()["fields"]
        assert stored.get("extra_subject") == "TEST_HIST_SUBJECT_XYZ"
        assert stored.get("extra_other") == "TEST_HIST_OTHER_ABC"
        assert stored.get("note") == "TEST_HIST_NOTE_QQQ"
        # Re-download and check markers present
        r3 = requests.get(f"{API}/contracts/{cid}/download", params={"format": "docx"})
        assert r3.status_code == 200
        with zipfile.ZipFile(io.BytesIO(r3.content)) as zf:
            xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
        assert "TEST_HIST_SUBJECT_XYZ" in xml
        assert "TEST_HIST_OTHER_ABC" in xml
        assert "TEST_HIST_NOTE_QQQ" in xml
        assert "Дополнительно: TEST_HIST_NOTE_QQQ" in xml
        assert "{{" not in xml and "{%p" not in xml
    finally:
        requests.delete(f"{API}/contracts/{cid}")


def test_existing_seeded_contract_present():
    """Seed contract from problem statement should be listed and downloadable."""
    seed_id = "45e95300-1ba1-4267-8f87-6b999ee5c83a"
    r = requests.get(f"{API}/contracts/{seed_id}")
    if r.status_code == 404:
        pytest.skip("Seeded contract not present in DB")
    assert r.status_code == 200
    data = r.json()
    assert "Шаназаров" in data.get("full_name", "") or "003370" in data.get("contract_number", "")
    # download docx
    r2 = requests.get(f"{API}/contracts/{seed_id}/download", params={"format": "docx"})
    assert r2.status_code == 200
    with zipfile.ZipFile(io.BytesIO(r2.content)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    assert "{{" not in xml and "}}" not in xml
