"""
Разбор Excel «заселение» и вспомогательные функции для раздела
«Распределение по этажам» (residents).

Структура исходного файла (лист «Лист1»):
    Блок ("902/2") | Ф.И.О | Статус | Группа | Дата заселения |
    Номер договора | Срок действия | Примечание
"""
import re
from datetime import datetime, date

# --- сопоставление заголовков колонок -> ключ ---------------------------------
_HEADER_MAP = {
    "блок": "block_room",
    "ф.и.о": "full_name",
    "фио": "full_name",
    "ф. и. о.": "full_name",
    "ф.и.о.": "full_name",
    "ф.и,о": "full_name",
    "статус": "status",
    "группа": "study_group",
    "дата заселения": "move_in_date",
    "дата вселения": "move_in_date",
    "номер договора": "contract_number",
    "№ договора": "contract_number",
    "договор": "contract_number",
    "срок действия": "term",
    "срок": "term",
    "примечание": "note",
    "примечания": "note",
}


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y")
    s = str(v).strip()
    # "2024-07-25 00:00:00" -> "25.07.2024"
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})(?:\s+00:00:00)?$", s)
    if m:
        return f"{m.group(3)}.{m.group(2)}.{m.group(1)}"
    return s


def parse_block_room(raw: str):
    """'902/2' -> (floor=9, block='902', room='2'). Возвращает (floor, block, room)."""
    s = (raw or "").strip()
    m = re.match(r"^\s*(\d{3,4})\s*/\s*(\d+)", s)
    if not m:
        # только блок без комнаты
        m2 = re.match(r"^\s*(\d{3,4})\s*$", s)
        if m2:
            block = m2.group(1)
            return _floor_of(block), block, ""
        return 0, "", ""
    block = m.group(1)
    room = m.group(2)
    return _floor_of(block), block, room


def _floor_of(block: str) -> int:
    if not block:
        return 0
    try:
        return int(block[:-2]) if len(block) >= 3 else int(block)
    except ValueError:
        return 0


def block_index(block: str) -> int:
    """'902' -> 2 (номер блока на этаже, 1..15)."""
    try:
        return int(block[-2:])
    except (ValueError, IndexError):
        return 0


def block_layout(rooms_present):
    """Стандартная планировка блока: малая комната /2 (2 места) + большая /4 (4 места).
    Если в данных встречается комната /3 — большая на 3 места. Возвращает [(room, capacity)]."""
    present = set(rooms_present or [])
    big = "3" if "3" in present else "4"
    layout = [("2", 2), (big, int(big))]
    # на случай нестандартных комнат — добавим их как есть
    for r in sorted(present):
        if r not in ("2", big):
            try:
                layout.append((r, int(r)))
            except ValueError:
                pass
    return layout


def block_capacity(rooms_present) -> int:
    return sum(cap for _, cap in block_layout(rooms_present))


def normalize_name(fio: str) -> str:
    """Нормализация ФИО для сопоставления с договором."""
    s = (fio or "").lower().strip()
    s = s.replace("ё", "е")
    s = re.sub(r"\s+", " ", s)
    return s


def _find_header_row(rows):
    """Ищем строку-заголовок, где встречается 'Блок' и 'Ф.И.О'."""
    for idx, r in enumerate(rows[:10]):
        cells = [str(c).strip().lower() if c is not None else "" for c in r]
        if any(c == "блок" for c in cells) and any(
            c in ("ф.и.о", "фио", "ф.и.о.", "ф. и. о.") for c in cells
        ):
            return idx, cells
    return None, None


def parse_zaselenie_xlsx(content: bytes):
    """Возвращает список словарей-жильцов из Excel заселения."""
    from openpyxl import load_workbook
    import io

    wb = load_workbook(io.BytesIO(content), data_only=True)

    best = None  # (people_count, list)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        hidx, header = _find_header_row(rows)
        if hidx is None:
            continue

        # карта: индекс колонки -> ключ
        col_key = {}
        for ci, h in enumerate(header):
            key = _HEADER_MAP.get((h or "").strip())
            if key and ci not in col_key.values():
                col_key[ci] = key
        # колонка примечания = сразу после «срок действия», если без заголовка
        term_ci = next((ci for ci, k in col_key.items() if k == "term"), None)
        if term_ci is not None and (term_ci + 1) not in col_key:
            col_key[term_ci + 1] = "note"

        people = []
        for r in rows[hidx + 1:]:
            rec = {"block_room": "", "full_name": "", "status": "",
                   "study_group": "", "move_in_date": "", "contract_number": "",
                   "term": "", "note": ""}
            for ci, key in col_key.items():
                if ci < len(r):
                    val = _clean(r[ci])
                    if val:
                        rec[key] = val
            if not rec["full_name"]:
                continue  # пустой слот в комнате — пропускаем
            floor, block, room = parse_block_room(rec["block_room"])
            if not block:
                continue
            people.append({
                "floor": floor,
                "block": block,
                "room": room,
                "full_name": rec["full_name"],
                "status": rec["status"],
                "study_group": rec["study_group"],
                "benefit": "",
                "contract_number": rec["contract_number"],
                "move_in_date": rec["move_in_date"],
                "term": rec["term"],
                "note": rec["note"],
            })
        if best is None or len(people) > best[0]:
            best = (len(people), people)

    return best[1] if best else []


# ---------------------------------------------------------------------------
# Разбор Word-списков групп: «ГРУППА СД-201» + список ФИО под ней
# ---------------------------------------------------------------------------
_GROUP_HEADER_RE = re.compile(r"^\s*группа\s+(.+?)\s*$", re.IGNORECASE)


def _looks_like_name(s: str) -> bool:
    s = (s or "").strip()
    if len(s) < 3:
        return False
    # хотя бы одна кириллическая/латинская буква, не строка-заголовок/номер
    if not re.search(r"[А-Яа-яЁёA-Za-z]", s):
        return False
    if s.isdigit():
        return False
    return True


def parse_group_lists_docx(content: bytes):
    """Возвращает список (fio, group) из Word-файла со списками групп.

    Формат: абзац «ГРУППА XXX» задаёт текущую группу, следующие абзацы — ФИО.
    Учитываются и таблицы (на случай, если списки оформлены таблицей)."""
    import io
    from docx import Document

    doc = Document(io.BytesIO(content))
    pairs = []
    current = None

    def handle_line(text):
        nonlocal current
        t = (text or "").strip()
        if not t:
            return
        m = _GROUP_HEADER_RE.match(t)
        if m:
            current = re.sub(r"\s+", " ", m.group(1).strip())
            return
        if current and _looks_like_name(t):
            # убираем ведущую нумерацию "1. ", "12) "
            t = re.sub(r"^\s*\d+[.)]\s*", "", t).strip()
            if _looks_like_name(t):
                pairs.append((t, current))

    for p in doc.paragraphs:
        handle_line(p.text)

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    handle_line(p.text)

    return pairs

