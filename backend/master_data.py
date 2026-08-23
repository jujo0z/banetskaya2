"""Единый Excel-шаблон «Данные» для всех документов.

Одна строка = один человек = все его документы (договор найма, сообщение,
Форма 19, Форма 24). Модуль умеет:
  * build_master_xlsx()      — сформировать стильный .xlsx с подсказками и
                               заполненной пробными данными первой строкой;
  * parse_master_xlsx(bytes) — прочитать заполненный файл в список словарей;
  * master_to_contract / _forma19 / _forma24 / _soobshenie — преобразовать
    строку-«человека» в поля конкретного документа.
"""
import io
import re

from document_service import _split_ru_date, _split_passport

# ---------------------------------------------------------------------------
# Схема единого шаблона. Каждая колонка: key / header / sample / section.
# dropdown — необязательный список допустимых кодов (выпадающий список в Excel).
# ---------------------------------------------------------------------------
SECTION_COLORS = {
    "Общие данные": "E11D48",
    "Паспорт": "7C3AED",
    "Место рождения": "0EA5E9",
    "Место жительства (регистрация)": "059669",
    "Откуда прибыл": "D97706",
    "Договор найма": "DB2777",
    "Регистрация / Сообщение": "0891B2",
    "Формы 19 / 24 — дополнительно": "4F46E5",
}

MASTER_COLUMNS = [
    # --- Общие данные ---
    {"key": "fio", "header": "ФИО (полностью)", "sample": "Иванов Иван Иванович", "section": "Общие данные", "width": 26},
    {"key": "birth_date", "header": "Дата рождения (ДД.ММ.ГГГГ)", "sample": "15.03.2004", "section": "Общие данные", "width": 16},
    {"key": "sex", "header": "Пол (1-муж, 2-жен)", "sample": "1", "section": "Общие данные", "width": 12, "dropdown": ["1", "2"]},
    {"key": "nationality", "header": "Национальность", "sample": "белорус", "section": "Общие данные", "width": 15},
    {"key": "citizenship", "header": "Гражданство", "sample": "Республика Беларусь", "section": "Общие данные", "width": 20},
    {"key": "phone", "header": "Телефон", "sample": "+375291234567", "section": "Общие данные", "width": 16},
    {"key": "id_number", "header": "Идентификационный номер (ИИН)", "sample": "4150304A011PB5", "section": "Общие данные", "width": 22},
    # --- Паспорт ---
    {"key": "passport", "header": "Паспорт (серия и номер)", "sample": "MP 1234567", "section": "Паспорт", "width": 18},
    {"key": "passport_issue_date", "header": "Паспорт: дата выдачи", "sample": "10.01.2020", "section": "Паспорт", "width": 16},
    {"key": "passport_valid_until", "header": "Паспорт: действителен до", "sample": "10.01.2030", "section": "Паспорт", "width": 18},
    {"key": "passport_issued_by", "header": "Паспорт: кем выдан", "sample": "Первомайским РУВД г. Минска", "section": "Паспорт", "width": 28},
    # --- Место рождения ---
    {"key": "bp_obl", "header": "Место рождения: область", "sample": "Минская", "section": "Место рождения", "width": 18},
    {"key": "bp_raion", "header": "Место рождения: район", "sample": "Минский", "section": "Место рождения", "width": 16},
    {"key": "bp_city", "header": "Место рождения: город (пгт)", "sample": "Минск", "section": "Место рождения", "width": 18},
    {"key": "bp_village", "header": "Место рождения: село (деревня)", "sample": "", "section": "Место рождения", "width": 18},
    # --- Место жительства (регистрация) ---
    {"key": "res_obl", "header": "Жительство: область", "sample": "Минская", "section": "Место жительства (регистрация)", "width": 16},
    {"key": "res_raion", "header": "Жительство: район", "sample": "Минский", "section": "Место жительства (регистрация)", "width": 15},
    {"key": "res_city", "header": "Жительство: город (пгт)", "sample": "Минск", "section": "Место жительства (регистрация)", "width": 16},
    {"key": "res_village", "header": "Жительство: село (деревня)", "sample": "", "section": "Место жительства (регистрация)", "width": 16},
    {"key": "res_street", "header": "Жительство: улица", "sample": "пр-т Дзержинского", "section": "Место жительства (регистрация)", "width": 20},
    {"key": "res_house", "header": "Жительство: дом", "sample": "85", "section": "Место жительства (регистрация)", "width": 10},
    {"key": "res_korpus", "header": "Жительство: корпус", "sample": "", "section": "Место жительства (регистрация)", "width": 12},
    {"key": "res_apartment", "header": "Жительство: комната/квартира", "sample": "302/2", "section": "Место жительства (регистрация)", "width": 18},
    # --- Откуда прибыл ---
    {"key": "from_obl", "header": "Откуда прибыл: область", "sample": "Гомельская", "section": "Откуда прибыл", "width": 18},
    {"key": "from_raion", "header": "Откуда прибыл: район", "sample": "Гомельский", "section": "Откуда прибыл", "width": 16},
    {"key": "from_city", "header": "Откуда прибыл: город (пгт)", "sample": "Гомель", "section": "Откуда прибыл", "width": 18},
    {"key": "from_village", "header": "Откуда прибыл: село (деревня)", "sample": "", "section": "Откуда прибыл", "width": 18},
    {"key": "arrival_date", "header": "Дата прибытия", "sample": "01.09.2024", "section": "Откуда прибыл", "width": 15},
    {"key": "lived_since", "header": "Проживал там с", "sample": "01.01.2010", "section": "Откуда прибыл", "width": 15},
    # --- Договор найма ---
    {"key": "contract_number", "header": "Номер договора", "sample": "0047390 003370", "section": "Договор найма", "width": 18},
    {"key": "sign_date", "header": "Дата подписания договора", "sample": "21.07.2025", "section": "Договор найма", "width": 18},
    {"key": "order_number", "header": "Номер приказа", "sample": "228", "section": "Договор найма", "width": 14},
    {"key": "order_date", "header": "Дата приказа", "sample": "21.07.2025", "section": "Договор найма", "width": 14},
    {"key": "room_number", "header": "Номер комнаты", "sample": "302/2", "section": "Договор найма", "width": 14},
    {"key": "contract_end_date", "header": "Срок договора до", "sample": "30.06.2028", "section": "Договор найма", "width": 16},
    # --- Регистрация / Сообщение ---
    {"key": "reg_organ", "header": "Орган регистрации", "sample": "Первомайский РУВД г. Минска", "section": "Регистрация / Сообщение", "width": 26},
    {"key": "reg_number", "header": "№ сообщения", "sample": "125", "section": "Регистрация / Сообщение", "width": 12},
    {"key": "reg_date", "header": "Дата сообщения", "sample": "01.09.2024", "section": "Регистрация / Сообщение", "width": 15},
    {"key": "reg_chief", "header": "Начальник (ФИО)", "sample": "Петров П.П.", "section": "Регистрация / Сообщение", "width": 18},
    {"key": "reg_from_date", "header": "Зарегистрирован с", "sample": "01.09.2024", "section": "Регистрация / Сообщение", "width": 17},
    {"key": "reg_to_date", "header": "Зарегистрирован по", "sample": "30.06.2028", "section": "Регистрация / Сообщение", "width": 17},
    # --- Формы 19 / 24 — дополнительно ---
    {"key": "purpose", "header": "Цель приезда (текст, Ф.19)", "sample": "на обучение", "section": "Формы 19 / 24 — дополнительно", "width": 20},
    {"key": "purpose_choice", "header": "Цель приезда Ф.24 (1-работа, 2-обучение)", "sample": "2", "section": "Формы 19 / 24 — дополнительно", "width": 22, "dropdown": ["1", "2"]},
    {"key": "purpose_term", "header": "Срок пребывания (по...)", "sample": "30.06.2028", "section": "Формы 19 / 24 — дополнительно", "width": 18},
    {"key": "prev_work", "header": "Где и кем работал (прежнее место)", "sample": "учащийся", "section": "Формы 19 / 24 — дополнительно", "width": 24},
    {"key": "education", "header": "Образование (1-7)", "sample": "2", "section": "Формы 19 / 24 — дополнительно", "width": 14, "dropdown": ["1", "2", "3", "4", "5", "6", "7"]},
    {"key": "marital", "header": "Семейное положение (1-4)", "sample": "2", "section": "Формы 19 / 24 — дополнительно", "width": 18, "dropdown": ["1", "2", "3", "4"]},
    {"key": "spouse_together", "header": "Прибыл с супругом (5-да, 6-нет)", "sample": "6", "section": "Формы 19 / 24 — дополнительно", "width": 20, "dropdown": ["5", "6"]},
    {"key": "children_count", "header": "Детей до 14 лет", "sample": "0", "section": "Формы 19 / 24 — дополнительно", "width": 14},
]

MASTER_KEYS = [c["key"] for c in MASTER_COLUMNS]
_HEADER_TO_KEY = {c["header"].strip().lower(): c["key"] for c in MASTER_COLUMNS}


# ---------------------------------------------------------------------------
# Вспомогательные разборщики дат
# ---------------------------------------------------------------------------
def _split_dmy(s: str):
    """'01.09.2024' | '2024-09-01' -> ('01', '09', '2024'). Иначе ('', '', '')."""
    s = (s or "").strip()
    if not s:
        return "", "", ""
    m = re.match(r"^\s*(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})\s*$", s)
    if m:
        return m.group(1).zfill(2), m.group(2).zfill(2), m.group(3)
    m = re.match(r"^\s*(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})\s*$", s)
    if m:
        return m.group(3).zfill(2), m.group(2).zfill(2), m.group(1)
    return "", "", ""


def _fio_parts(fio: str):
    parts = (fio or "").split()
    sur = parts[0] if len(parts) >= 1 else ""
    name = parts[1] if len(parts) >= 2 else ""
    patr = " ".join(parts[2:]) if len(parts) >= 3 else ""
    return sur, name, patr


def _nz(*vals):
    """Первое непустое значение."""
    for v in vals:
        if v and str(v).strip():
            return str(v).strip()
    return ""


def _join(sep, *vals):
    return sep.join([str(v).strip() for v in vals if v and str(v).strip()])


# ---------------------------------------------------------------------------
# Преобразование строки-«человека» -> поля документов
# ---------------------------------------------------------------------------
def master_to_contract(m: dict) -> dict:
    m = m or {}
    address = _join(", ", m.get("res_street"), m.get("res_house"))
    if m.get("res_apartment"):
        address = _join(", ", address, f"ком. {m['res_apartment']}")
    return {
        "contract_number": m.get("contract_number", ""),
        "sign_date": m.get("sign_date", ""),
        "order_number": m.get("order_number", ""),
        "order_date": m.get("order_date", ""),
        "citizenship": m.get("citizenship", ""),
        "full_name": m.get("fio", ""),
        "birth_date": m.get("birth_date", ""),
        "room_number": m.get("room_number", ""),
        "contract_end_date": m.get("contract_end_date", ""),
        "registration_address": address,
        "passport_number": m.get("passport", ""),
        "passport_issue_date": m.get("passport_issue_date", ""),
        "passport_valid_until": m.get("passport_valid_until", ""),
        "passport_issued_by": m.get("passport_issued_by", ""),
        "id_number": m.get("id_number", ""),
        "phone": m.get("phone", ""),
    }


def master_to_forma24(m: dict) -> dict:
    m = m or {}
    sur, name, patr = _fio_parts(m.get("fio", ""))
    d, mo, y = _split_ru_date(m.get("birth_date", ""))
    ser, num = _split_passport(m.get("passport", ""))
    rec = {
        "surname": sur, "first_name": name, "patronymic": patr,
        "birth_day": d, "birth_month": mo, "birth_year": y,
        "sex": m.get("sex", ""),
        "nationality": m.get("nationality", ""),
        "citizenship": m.get("citizenship", ""),
        "bp_obl": m.get("bp_obl", ""), "bp_raion": m.get("bp_raion", ""),
        "bp_city": m.get("bp_city", ""), "bp_village": m.get("bp_village", ""),
        "res_obl": m.get("res_obl", ""), "res_raion": m.get("res_raion", ""),
        "res_city": m.get("res_city", ""), "res_village": m.get("res_village", ""),
        "from_obl": m.get("from_obl", ""), "from_raion": m.get("from_raion", ""),
        "from_city": m.get("from_city", ""), "from_village": m.get("from_village", ""),
        "arrival_date": m.get("arrival_date", ""),
        "lived_since": m.get("lived_since", ""),
        "purpose_choice": m.get("purpose_choice", ""),
        "term": m.get("purpose_term", ""),
        "prev_work": m.get("prev_work", ""),
        "education": m.get("education", ""),
        "marital": m.get("marital", ""),
        "spouse_together": m.get("spouse_together", ""),
        "children_count": m.get("children_count", ""),
    }
    return {k: v for k, v in rec.items() if str(v).strip()}


def master_to_forma19(m: dict) -> dict:
    m = m or {}
    sur, name, patr = _fio_parts(m.get("fio", ""))
    d, mo, y = _split_ru_date(m.get("birth_date", ""))
    ser, num = _split_passport(m.get("passport", ""))
    sx = str(m.get("sex", "")).strip()
    sx = "М" if sx == "1" else ("Ж" if sx == "2" else sx)
    rec = {
        "id_number": m.get("id_number", ""),
        "surname": sur, "first_name": name, "patronymic": patr,
        "birth_day": d, "birth_month": mo, "birth_year": y,
        "sex": sx,
        "citizenship": m.get("citizenship", ""),
        "bp_obl": m.get("bp_obl", ""), "bp_raion": m.get("bp_raion", ""),
        "bp_city": m.get("bp_city", ""), "bp_village": m.get("bp_village", ""),
        "res_obl": m.get("res_obl", ""), "res_raion": m.get("res_raion", ""),
        "res_city": m.get("res_city", ""), "res_village": m.get("res_village", ""),
        "res_street": m.get("res_street", ""), "res_house": m.get("res_house", ""),
        "res_korpus": m.get("res_korpus", ""), "res_apartment": m.get("res_apartment", ""),
        "reg_authority": _nz(m.get("reg_organ")),
        "from_obl": m.get("from_obl", ""), "from_raion": m.get("from_raion", ""),
        "from_city": m.get("from_city", ""), "from_village": m.get("from_village", ""),
        "arrival_date": m.get("arrival_date", ""),
        "purpose": m.get("purpose", ""),
        "purpose_term": m.get("purpose_term", ""),
        "employment": m.get("prev_work", ""),
        "passport_series": ser, "passport_number": num,
        "passport_issued": m.get("passport_issued_by", ""),
    }
    return {k: v for k, v in rec.items() if str(v).strip()}


def master_to_soobshenie(m: dict) -> dict:
    m = m or {}
    dd, dm, dy = _split_dmy(m.get("reg_date", ""))
    idd, idm, idy = _split_dmy(m.get("passport_issue_date", ""))
    fd, fm, fy = _split_dmy(m.get("reg_from_date", ""))
    td, tm, ty = _split_dmy(m.get("reg_to_date", ""))
    ser, num = _split_passport(m.get("passport", ""))
    _, _, by = _split_dmy(m.get("birth_date", ""))
    place = _nz(m.get("bp_city"), m.get("bp_village"), m.get("bp_raion"), m.get("bp_obl"))
    birth = _join(", ", (f"{by} г." if by else ""), place)
    address = _join(", ", m.get("res_city"), m.get("res_street"),
                    (f"д. {m['res_house']}" if m.get("res_house") else ""),
                    (f"ком. {m['res_apartment']}" if m.get("res_apartment") else ""))
    rec = {
        "reg_organ": m.get("reg_organ", ""),
        "date_day": dd, "date_month": dm, "date_year": dy[-2:] if dy else "",
        "number": m.get("reg_number", ""),
        "fio": m.get("fio", ""),
        "birth": birth,
        "address": address,
        "passport_series": ser, "passport_number": num,
        "issue_day": idd, "issue_month": idm, "issue_year": idy[-2:] if idy else "",
        "issued_by": m.get("passport_issued_by", ""),
        "from_day": fd, "from_month": fm, "from_year": fy[-2:] if fy else "",
        "to_day": td, "to_month": tm, "to_year": ty[-2:] if ty else "",
        "chief": m.get("reg_chief", ""),
    }
    return {k: v for k, v in rec.items() if str(v).strip()}


def derive_all(m: dict) -> dict:
    return {
        "contract": master_to_contract(m),
        "forma19": master_to_forma19(m),
        "forma24": master_to_forma24(m),
        "soobshenie": master_to_soobshenie(m),
    }


def contract_to_master(fields: dict) -> dict:
    """Обратное преобразование для старых договоров без сохранённого master."""
    f = fields or {}
    return {
        "fio": f.get("full_name", ""),
        "birth_date": f.get("birth_date", ""),
        "citizenship": f.get("citizenship", ""),
        "phone": f.get("phone", ""),
        "id_number": f.get("id_number", ""),
        "passport": f.get("passport_number", ""),
        "passport_issue_date": f.get("passport_issue_date", ""),
        "passport_valid_until": f.get("passport_valid_until", ""),
        "passport_issued_by": f.get("passport_issued_by", ""),
        "res_street": f.get("registration_address", ""),
        "contract_number": f.get("contract_number", ""),
        "sign_date": f.get("sign_date", ""),
        "order_number": f.get("order_number", ""),
        "order_date": f.get("order_date", ""),
        "room_number": f.get("room_number", ""),
        "contract_end_date": f.get("contract_end_date", ""),
    }


# ---------------------------------------------------------------------------
# Excel: генерация красивого шаблона
# ---------------------------------------------------------------------------
def build_master_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"

    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    white_bold = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    hdr_font = Font(name="Calibri", size=9, bold=True, color="111827")
    sample_font = Font(name="Calibri", size=10, color="1E3A8A", italic=True)
    hdr_fill = PatternFill("solid", fgColor="F3F4F6")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Строка 1 — секции (объединяем по группам), строка 2 — заголовки колонок.
    col = 1
    n = len(MASTER_COLUMNS)
    i = 0
    while i < n:
        sec = MASTER_COLUMNS[i]["section"]
        j = i
        while j < n and MASTER_COLUMNS[j]["section"] == sec:
            j += 1
        span = j - i
        c1 = get_column_letter(col)
        c2 = get_column_letter(col + span - 1)
        ws.merge_cells(f"{c1}1:{c2}1")
        cell = ws[f"{c1}1"]
        cell.value = sec
        cell.fill = PatternFill("solid", fgColor=SECTION_COLORS.get(sec, "6B7280"))
        cell.font = white_bold
        cell.alignment = center
        for k in range(span):
            ws[f"{get_column_letter(col + k)}1"].border = border
        col += span
        i = j

    # Заголовки колонок (строка 2) + пробные данные (строка 3)
    for idx, c in enumerate(MASTER_COLUMNS, start=1):
        L = get_column_letter(idx)
        h = ws[f"{L}2"]
        h.value = c["header"]
        h.font = hdr_font
        h.fill = hdr_fill
        h.alignment = center
        h.border = border
        s = ws[f"{L}3"]
        s.value = c["sample"]
        s.font = sample_font
        s.alignment = left
        s.border = border
        ws.column_dimensions[L].width = c.get("width", 16)

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 42
    ws.freeze_panes = "A3"

    # Выпадающие списки для кодовых колонок (строки 3..600)
    for idx, c in enumerate(MASTER_COLUMNS, start=1):
        if c.get("dropdown"):
            L = get_column_letter(idx)
            dv = DataValidation(type="list",
                                formula1='"%s"' % ",".join(c["dropdown"]),
                                allow_blank=True, showDropDown=False)
            dv.error = "Выберите значение из списка"
            dv.prompt = "Выберите код из списка"
            ws.add_data_validation(dv)
            dv.add(f"{L}3:{L}600")

    # --- Лист «Инструкция» ---
    ins = wb.create_sheet("Инструкция")
    ins.column_dimensions["A"].width = 34
    ins.column_dimensions["B"].width = 80
    title_fill = PatternFill("solid", fgColor="E11D48")
    ins.merge_cells("A1:B1")
    t = ins["A1"]
    t.value = "Как заполнять таблицу"
    t.font = Font(size=13, bold=True, color="FFFFFF")
    t.fill = title_fill
    t.alignment = Alignment(horizontal="left", vertical="center")
    ins.row_dimensions[1].height = 26

    rows = [
        ("Одна строка — один человек", "Заполните по одной строке на каждого человека, начиная со строки 3. Из этих данных формируются все документы: договор найма, сообщение, Форма 19 и Форма 24."),
        ("Формат дат", "Пишите даты как ДД.ММ.ГГГГ, например 15.03.2004."),
        ("Пол", "1 — мужской, 2 — женский."),
        ("Цель приезда (Форма 24)", "1 — на работу, 2 — на обучение."),
        ("Образование (код)", "1 — высшее; 2 — среднее специальное; 3 — профессионально-техническое; 4 — общее среднее; 5 — общее базовое; 6 — общее начальное; 7 — не имеет начального."),
        ("Семейное положение (код)", "1 — состоит в браке; 2 — никогда не состоял(а) в браке; 3 — вдовец (вдова); 4 — разведён(а)."),
        ("Прибыл с супругом", "5 — да, 6 — нет."),
        ("Паспорт", "Пишите серию и номер вместе, например «MP 1234567» — программа сама разделит их где нужно."),
        ("Пустые поля", "Ненужные для конкретного документа поля можно оставить пустыми."),
    ]
    r = 3
    lbl_font = Font(bold=True, color="111827")
    val_font = Font(color="374151")
    wrap = Alignment(vertical="top", wrap_text=True)
    for a, b in rows:
        ins[f"A{r}"].value = a
        ins[f"A{r}"].font = lbl_font
        ins[f"A{r}"].alignment = wrap
        ins[f"B{r}"].value = b
        ins[f"B{r}"].font = val_font
        ins[f"B{r}"].alignment = wrap
        ins.row_dimensions[r].height = 30
        r += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel: чтение заполненного шаблона
# ---------------------------------------------------------------------------
def _cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def has_master_headers(content: bytes) -> bool:
    """Похож ли файл на единый шаблон «Данные» (по распознанным заголовкам)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        return False
    ws = wb["Данные"] if "Данные" in wb.sheetnames else wb.active
    for row in list(ws.iter_rows(values_only=True))[:5]:
        hits = sum(1 for c in row if _cell_str(c).lower() in _HEADER_TO_KEY)
        if hits >= 3:
            return True
    return False


def parse_master_xlsx(content: bytes) -> list:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb["Данные"] if "Данные" in wb.sheetnames else wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return []

    # Найти строку заголовков (там, где встречаются известные заголовки колонок).
    header_idx = None
    for i, row in enumerate(all_rows[:5]):
        hits = 0
        for cell in row:
            if _cell_str(cell).lower() in _HEADER_TO_KEY:
                hits += 1
        if hits >= 3:
            header_idx = i
            break
    if header_idx is None:
        header_idx = 1 if len(all_rows) > 1 else 0

    header = all_rows[header_idx]
    col_key = {}
    for ci, cell in enumerate(header):
        k = _HEADER_TO_KEY.get(_cell_str(cell).lower())
        if k:
            col_key[ci] = k

    # Фолбэк: если заголовки не распознаны — сопоставляем по позиции.
    if not col_key:
        for ci, c in enumerate(MASTER_COLUMNS):
            col_key[ci] = c["key"]

    people = []
    for row in all_rows[header_idx + 1:]:
        if row is None or all(c is None for c in row):
            continue
        rec = {}
        for ci, key in col_key.items():
            if ci < len(row):
                rec[key] = _cell_str(row[ci])
        if any(v for v in rec.values()):
            people.append(rec)
    return people
