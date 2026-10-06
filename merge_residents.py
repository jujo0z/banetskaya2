#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Слияние нового файла заселения с текущей БД residents.
Запуск:
  python3 merge_residents.py            # DRY RUN (только отчёт, без записи)
  python3 merge_residents.py apply      # применить изменения
Безопасность: НЕ удаляет проживающих, которых нет в файле — только показывает списком.
Сохраняет уже имеющиеся номера договоров и сроки (берёт из файла только если там заполнено).
"""
import os, re, sys, io, uuid
from datetime import datetime, timezone, date
from difflib import SequenceMatcher
from openpyxl import load_workbook
from pymongo import MongoClient

XLSX = "/root/dlya_ivana.xlsx"
MONGO_URL = "mongodb://127.0.0.1:27017"
DB_NAME = "banetskaya_db"
APPLY = len(sys.argv) > 1 and sys.argv[1] == "apply"

GROUP_RE = re.compile(r"^[А-Яа-яЁё]{1,4}-\d{1,3}$")

def now_iso():
    return datetime.now(timezone.utc).isoformat()

def _clean(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y")
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})(?:\s+00:00:00)?$", s)
    if m:
        return f"{m.group(3)}.{m.group(2)}.{m.group(1)}"
    return s

def _floor_of(block):
    if not block:
        return 0
    try:
        return int(block[:-2]) if len(block) >= 3 else int(block)
    except ValueError:
        return 0

def parse_block_room(raw):
    s = (raw or "").strip()
    m = re.match(r"^\s*(\d{3,4})\s*/\s*(\d+)", s)
    if not m:
        m2 = re.match(r"^\s*(\d{3,4})\s*$", s)
        if m2:
            b = m2.group(1)
            return _floor_of(b), b, ""
        return 0, "", ""
    b = m.group(1)
    return _floor_of(b), b, m.group(2)

def normalize_name(fio):
    s = (fio or "").lower().strip()
    s = s.replace("ё", "е")
    s = re.sub(r"\s+", " ", s)
    return s

# ---- parse file ----
wb = load_workbook(XLSX, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
# header row index: ищем 'блок' и 'ф.и.о'
hidx = None
for i, r in enumerate(rows[:10]):
    cells = [str(c).strip().lower() if c is not None else "" for c in r]
    if any(c == "блок" for c in cells) and any(c in ("ф.и.о","фио","ф.и.о.") for c in cells):
        hidx = i
        break
assert hidx is not None, "header not found"
# фикс. индексы колонок по файлу: 0 №, 1 Блок, 2 ФИО, 3 Статус, 4 Группа, 5 Дата, 6 Номер дог, 7 Срок
people = []
shift_fixed = 0
for r in rows[hidx+1:]:
    def g(i):
        return _clean(r[i]) if i < len(r) else ""
    fio = g(2)
    if not fio:
        continue
    floor, block, room = parse_block_room(g(1))
    if not block:
        continue
    status = g(3)
    group = g(4)
    move_in = g(5)
    contract = g(6)
    term = g(7)
    note = g(8)
    # фикс смещения: если группа пустая, а в 'срок' стоит код группы -> вернуть в группу
    if not group and term and GROUP_RE.match(term):
        group = term
        term = ""
        shift_fixed += 1
    people.append({
        "floor": floor, "block": block, "room": room,
        "full_name": fio, "status": status, "study_group": group,
        "move_in_date": move_in, "contract_number": contract,
        "term": term, "note": note,
    })

# ---- load DB ----
cli = MongoClient(MONGO_URL)
db = cli[DB_NAME]
db_docs = list(db.residents.find({}))
by_norm = {}
for d in db_docs:
    by_norm.setdefault(normalize_name(d.get("full_name","")), []).append(d)

def toks(name):
    return normalize_name(name).split()

def build_update(d, p):
    new_contract = p["contract_number"] or d.get("contract_number","")
    new_term = p["term"] or d.get("term","")
    preserved = (not p["contract_number"]) and bool(d.get("contract_number"))
    vals = {
        "floor": p["floor"], "block": p["block"], "room": p["room"],
        "full_name": p["full_name"],
        "status": p["status"] or d.get("status",""),
        "study_group": p["study_group"] or d.get("study_group",""),
        "move_in_date": p["move_in_date"] or d.get("move_in_date",""),
        "contract_number": new_contract,
        "term": new_term,
        "updated_at": now_iso(),
    }
    return vals, preserved

used_ids = set()
seen_file = {}
updated, inserted, moved, file_dupes, name_fixed = [], [], [], [], []
preserved_contract = 0
to_insert_docs = []
to_update_ops = []
pending = []  # не совпали точно -> во второй проход

# --- ПРОХОД 1: точное совпадение по нормализованному ФИО ---
for p in people:
    norm = normalize_name(p["full_name"])
    if norm in seen_file:
        file_dupes.append(p["full_name"] + " -> " + p["block"] + "/" + p["room"])
        continue
    seen_file[norm] = True
    cands = [d for d in by_norm.get(norm, []) if d["_id"] not in used_ids]
    if cands:
        d = cands[0]
        used_ids.add(d["_id"])
        vals, preserved = build_update(d, p)
        if preserved: preserved_contract += 1
        if (d.get("floor"),d.get("block"),d.get("room")) != (p["floor"],p["block"],p["room"]):
            moved.append(f'{p["full_name"]}: {d.get("block")}/{d.get("room")} -> {p["block"]}/{p["room"]}')
        to_update_ops.append((d["_id"], vals))
        updated.append(p["full_name"])
    else:
        pending.append(p)

# --- ПРОХОД 2: нечёткое сопоставление (исправленная орфография/добавленное отчество) ---
def fuzzy_match(p, free_docs):
    pt = toks(p["full_name"])
    p_surname = pt[0] if pt else ""
    p_first = pt[1] if len(pt) > 1 else ""
    best, best_score = None, 0.0
    for d in free_docs:
        dt = toks(d.get("full_name",""))
        d_surname = dt[0] if dt else ""
        d_first = dt[1] if len(dt) > 1 else ""
        full_ratio = SequenceMatcher(None, normalize_name(p["full_name"]), normalize_name(d.get("full_name",""))).ratio()
        sur_ratio = SequenceMatcher(None, p_surname, d_surname).ratio()
        score = 0.0
        # тот же блок + то же имя + похожая фамилия -> почти наверняка один человек
        if p["block"] == d.get("block") and p_first and p_first == d_first and sur_ratio >= 0.6:
            score = 1.5 + sur_ratio
        # очень высокое сходство полного ФИО
        elif full_ratio >= 0.88:
            score = full_ratio
        if score > best_score:
            best, best_score = d, score
    return best

for p in pending:
    free_docs = [d for d in db_docs if d["_id"] not in used_ids]
    d = fuzzy_match(p, free_docs)
    if d is not None:
        used_ids.add(d["_id"])
        vals, preserved = build_update(d, p)
        if preserved: preserved_contract += 1
        if d.get("full_name") != p["full_name"]:
            name_fixed.append(f'{d.get("full_name")}  ->  {p["full_name"]}')
        if (d.get("floor"),d.get("block"),d.get("room")) != (p["floor"],p["block"],p["room"]):
            moved.append(f'{p["full_name"]}: {d.get("block")}/{d.get("room")} -> {p["block"]}/{p["room"]}')
        to_update_ops.append((d["_id"], vals))
        updated.append(p["full_name"])
    else:
        doc = {
            "id": str(uuid.uuid4()),
            "floor": p["floor"], "block": p["block"], "room": p["room"],
            "full_name": p["full_name"], "status": p["status"],
            "study_group": p["study_group"], "benefit": "",
            "contract_number": p["contract_number"], "move_in_date": p["move_in_date"],
            "term": p["term"], "note": p["note"],
            "created_at": now_iso(), "updated_at": now_iso(),
            "no_contract_needed": False,
        }
        to_insert_docs.append(doc)
        inserted.append(p["full_name"])

removed = [d for d in db_docs if d["_id"] not in used_ids]

print("="*60)
print("MODE:", "APPLY (запись в БД)" if APPLY else "DRY RUN (без записи)")
print("="*60)
print(f"Строк-студентов в файле:        {len(people)}")
print(f"Исправлено смещений колонок:    {shift_fixed}")
print(f"Проживающих сейчас в БД:        {len(db_docs)}")
print(f"Дубли ФИО внутри файла:         {len(file_dupes)}")
print("-"*60)
print(f"Обновлено (совпали по ФИО):     {len(updated)}")
print(f"  из них исправлено имя:        {len(name_fixed)}")
print(f"  из них сохранён номер дог.:   {preserved_contract}")
print(f"Новых (добавить):               {len(inserted)}")
print(f"Переехали (блок/комната):       {len(moved)}")
print(f"Нет в файле (НЕ удаляем):       {len(removed)}")
print("="*60)
print("\n--- ДУБЛИ ВНУТРИ ФАЙЛА (пропущены вторые вхождения) ---")
for x in file_dupes: print(" ", x)
print("\n--- ИСПРАВЛЕНЫ ИМЕНА (старое -> новое из файла) ---")
for x in name_fixed: print(" ", x)
print("\n--- НЕТ В НОВОМ ФАЙЛЕ (останутся в базе, помечены) ---")
for d in removed: print(f'  {d.get("block")}/{d.get("room")}  {d.get("full_name")}  [дог:{d.get("contract_number") or "-"}]')
print("\n--- НОВЫЕ (будут добавлены) [первые 60] ---")
for x in inserted[:60]: print(" ", x)
if len(inserted) > 60: print(f"  ... ещё {len(inserted)-60}")
print("\n--- ПЕРЕЕХАЛИ [первые 40] ---")
for x in moved[:40]: print(" ", x)
if len(moved) > 40: print(f"  ... ещё {len(moved)-40}")

if APPLY:
    print("\n>>> ПРИМЕНЯЮ ИЗМЕНЕНИЯ...")
    for _id, vals in to_update_ops:
        db.residents.update_one({"_id": _id}, {"$set": vals})
    if to_insert_docs:
        db.residents.insert_many(to_insert_docs)
    # пометим тех, кого нет в файле
    rem_ids = [d["_id"] for d in removed]
    if rem_ids:
        db.residents.update_many({"_id": {"$in": rem_ids}},
            {"$set": {"not_in_list_2026_10": True, "updated_at": now_iso()}})
    print("Готово. residents в БД теперь:", db.residents.count_documents({}))
    # распределение по этажам
    print("--- этажи после ---")
    for row in db.residents.aggregate([{"$group":{"_id":"$floor","n":{"$sum":1}}},{"$sort":{"_id":1}}]):
        print("  этаж", row["_id"], ":", row["n"])
else:
    print("\n(DRY RUN — ничего не записано. Для применения: python3 merge_residents.py apply)")
