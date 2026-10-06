#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сверка по НОМЕРУ договора: file <-> residents <-> contracts."""
import re
from datetime import datetime, date
from openpyxl import load_workbook
from pymongo import MongoClient
db=MongoClient("mongodb://127.0.0.1:27017")["banetskaya_db"]
def norm(f): return re.sub(r"\s+"," ",(f or "").lower().replace("ё","е").strip())
def dig(s):
    d=re.sub(r"\D","",s or "")
    return d.lstrip("0") or ("0" if d else "")
def cl(v):
    if v is None: return ""
    if isinstance(v,(datetime,date)): return v.strftime("%d.%m.%Y")
    s=str(v).strip(); m=re.match(r"^(\d{4})-(\d{2})-(\d{2})",s)
    return f"{m.group(3)}.{m.group(2)}.{m.group(1)}" if m else s

wb=load_workbook("/root/dlya_ivana.xlsx",data_only=True); ws=wb[wb.sheetnames[0]]
rows=list(ws.iter_rows(values_only=True))
file_by_num={}; file_people=[]; seen=set()
for r in rows[3:]:
    def g(i): return cl(r[i]) if i<len(r) else ""
    if not g(2): continue
    nk=norm(g(2))
    if nk in seen: continue
    seen.add(nk)
    num=dig(g(6))
    p={"name":g(2),"num":num,"block":g(1)}
    file_people.append(p)
    if num: file_by_num.setdefault(num,p)
file_nums=set(file_by_num)
file_names=set(norm(p["name"]) for p in file_people)

res=list(db.residents.find({})); con=list(db.contracts.find({}))
def innum(d): return dig(d.get("contract_number",""))

# выбывшие = нет ни по номеру, ни по ФИО в файле
res_out=[d for d in res if innum(d) not in file_nums and norm(d.get("full_name","")) not in file_names]
con_out=[d for d in con if innum(d) not in file_nums and norm(d.get("full_name","")) not in file_names]

# пары разного ФИО: по номеру договора file<->contract
con_by_num={}
for d in con:
    n=innum(d)
    if n: con_by_num.setdefault(n,[]).append(d)
name_pairs=[]
for num,p in file_by_num.items():
    cs=con_by_num.get(num)
    if cs:
        c=cs[0]
        if norm(p["name"])!=norm(c.get("full_name","")):
            name_pairs.append((num,p["name"],c.get("full_name")))

# файловые люди без договора (по номеру и по имени)
con_names=set(norm(d.get("full_name","")) for d in con)
con_nums=set(innum(d) for d in con if innum(d))
file_no_contract=[p for p in file_people if (not p["num"] or p["num"] not in con_nums) and norm(p["name"]) not in con_names]

print("Файл (уник.людей):",len(file_people)," с номером:",len(file_nums))
print("residents:",len(res)," contracts:",len(con))
print("ВЫБЫВШИЕ residents (удалить):",len(res_out))
print("ВЫБЫВШИЕ contracts (удалить):",len(con_out))
print("Файловые БЕЗ договора (создать):",len(file_no_contract))
print("ПАРЫ разного ФИО (один №договора):",len(name_pairs))
print("\n=== ВЫБЫВШИЕ residents ===")
for d in res_out: print(f'  {d.get("block")}/{d.get("room")}  {d.get("full_name")}  №{d.get("contract_number")}')
print("\n=== ВЫБЫВШИЕ contracts ===")
for d in con_out: print(f'  {d.get("full_name")}  №{d.get("contract_number")}  status={d.get("status")}')
print("\n=== ФАЙЛОВЫЕ БЕЗ ДОГОВОРА ===")
for p in file_no_contract: print(f'  {p["block"]}  {p["name"]}  №{p["num"]}')
print("\n=== ВСЕ ПАРЫ РАЗНОГО ФИО (№ | файл/заселение | карточка договора) ===")
for i,(num,fn,cn) in enumerate(sorted(name_pairs),1):
    print(f'  {i:2}. №{num}  ФАЙЛ: {fn}   |   ДОГОВОР: {cn}')
